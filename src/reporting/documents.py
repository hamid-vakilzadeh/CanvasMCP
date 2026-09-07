"""Bounded, read-only extraction with source locations and explicit coverage gaps.

Parser references were checked with Context7 on 2026-09-07. No OCR, formula
calculation, macros, external links, or embedded programs are executed.
"""

from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import json
import os
import re
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import zipfile
from urllib.parse import urljoin, urlsplit

import httpx2

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".txt", ".md", ".html", ".htm", ".json", ".py", ".r", ".sql"}


def extract_document(data: bytes, filename: str) -> dict:
    suffix = Path(filename).suffix.lower()
    segments: list[dict] = []
    gaps: list[dict] = []
    char_count = 0

    def add(location, text):
        nonlocal char_count
        text = str(text if text is not None else "").strip()
        if not text:
            return
        char_count += len(text)
        if char_count > MAX_TEXT_CHARS:
            raise ValueError("document_text_limit")
        segments.append({"location": location, "text": text})

    def gap(location, reason):
        gaps.append({"location": location, "reason": reason})

    if suffix not in SUPPORTED:
        return {"segments": [], "gaps": [{"location": "file", "reason": "unsupported_format"}], "complete": False}
    if len(data) > MAX_FILE_BYTES:
        return {"segments": [], "gaps": [{"location": "file", "reason": "file_size_limit"}], "complete": False}
    try:
        if suffix in {".docx", ".pptx", ".xlsx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES or len(entries) > 10_000:
                    raise ValueError("document_expansion_limit")
                names = [e.filename for e in entries]
                if any('/media/' in name or '/embeddings/' in name or '/charts/' in name for name in names):
                    gap("embedded content", "visual_or_embedded_content_not_reviewed")
                if any(name.endswith('vbaProject.bin') for name in names):
                    gap("macros", "macros_not_executed")
        if suffix == ".pdf":
            from pypdf import PdfReader
            document = PdfReader(io.BytesIO(data))
            if document.is_encrypted and not document.decrypt(""):
                raise ValueError("encrypted_document")
            if len(document.pages) > 2000:
                raise ValueError("document_page_limit")
            for index, page in enumerate(document.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    add(f"page {index}", text)
                else:
                    gap(f"page {index}", "no_extractable_text_or_scanned_page")
                resources = page.get('/Resources')
                if resources is not None and hasattr(resources, 'get_object'):
                    resources = resources.get_object()
                if resources and resources.get('/XObject'):
                    gap(f"page {index}", "visual_content_not_reviewed")
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table
            document = Document(io.BytesIO(data))

            def walk(container, prefix):
                for index, element in enumerate(container.iter_inner_content(), 1):
                    location = f"{prefix} block {index}"
                    if isinstance(element, Table):
                        for r, row in enumerate(element.rows, 1):
                            for c, cell in enumerate(row.cells, 1):
                                walk(cell, f"{location}, row {r} cell {c}")
                    else:
                        add(location, element.text)
            walk(document, "document")
            for i, section in enumerate(document.sections, 1):
                for name in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
                    part = getattr(section, name)
                    if not part.is_linked_to_previous:
                        walk(part, f"section {i} {name}")
            # python-docx does not expose all text in text boxes, comments and footnotes.
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                from lxml import etree
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for name in ('word/footnotes.xml', 'word/endnotes.xml', 'word/comments.xml'):
                    if name in archive.namelist():
                        xml = etree.fromstring(archive.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                        for i, item in enumerate(xml, 1):
                            add(f"{Path(name).stem} {i}", '\n'.join(item.xpath('.//w:t/text()', namespaces=ns)))
                xml = etree.fromstring(archive.read('word/document.xml'), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for i, box in enumerate(xml.xpath('.//w:txbxContent', namespaces=ns), 1):
                    add(f"text box {i}", '\n'.join(box.xpath('.//w:t/text()', namespaces=ns)))
        elif suffix == ".pptx":
            from pptx import Presentation
            document = Presentation(io.BytesIO(data))

            def shapes(values, prefix):
                for i, shape in enumerate(values, 1):
                    where = f"{prefix}, shape {i}"
                    if hasattr(shape, 'shapes'):
                        shapes(shape.shapes, where)
                    if shape.has_text_frame:
                        add(where, shape.text_frame.text)
                    if shape.has_table:
                        for r, row in enumerate(shape.table.rows, 1):
                            add(f"{where}, row {r}", ' | '.join(c.text for c in row.cells))
            for i, slide in enumerate(document.slides, 1):
                shapes(slide.shapes, f"slide {i}")
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    add(f"slide {i} speaker notes", slide.notes_slide.notes_text_frame.text)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in formulas:
                    cached_sheet = values[sheet.title]
                    # Ignore a corrupt producer's dimension declaration; scan the real XML.
                    sheet.reset_dimensions()
                    cached_sheet.reset_dimensions()
                    for row, cached_row in zip(sheet.iter_rows(), cached_sheet.iter_rows(), strict=True):
                        for cell, cached in zip(row, cached_row, strict=True):
                            if cell.value is None:
                                continue
                            where = f"sheet {sheet.title}, cell {cell.coordinate}"
                            if cell.data_type == 'f':
                                add(where, f"Formula: {cell.value}\nCached result: {cached.value if cached.value is not None else 'unavailable'}")
                                if cached.value is None:
                                    gap(where, "formula_result_unavailable_not_calculated")
                            else:
                                add(where, cell.value)
            finally:
                formulas.close()
                values.close()
        else:
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                if data.startswith((b'\xff\xfe', b'\xfe\xff')):
                    text = data.decode('utf-16')
                else:
                    raise ValueError('unsupported_text_encoding')
            if suffix in {'.csv', '.tsv'}:
                for i, row in enumerate(csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','), 1):
                    add(f"row {i}", ' | '.join(row))
            elif suffix in {'.html', '.htm'}:
                from reporting.common import plain
                add('document', plain(text))
                if re.search(r'<(?:img|svg|canvas|audio|video|iframe|object)\b', text, re.I):
                    gap('embedded content', 'visual_or_embedded_content_not_reviewed')
            else:
                all_lines = text.splitlines()
                for start in range(0, len(all_lines), 100):
                    lines = all_lines[start:start + 100]
                    add(f"lines {start + 1}–{start + len(lines)}", '\n'.join(lines))
        if not segments and not gaps:
            gap("file", "no_extractable_content")
    except Exception as exc:
        known = str(exc) if isinstance(exc, ValueError) and str(exc) in {
            'document_text_limit', 'document_expansion_limit', 'document_page_limit',
            'encrypted_document', 'unsupported_text_encoding'} else 'document_parse_failed'
        gap("file", known)
    return {"segments": segments, "gaps": gaps, "complete": not gaps, "characters": sum(len(s['text']) for s in segments)}


async def extract_isolated(data: bytes, filename: str, private_directory: Path) -> dict:
    """Run parsers without Canvas credentials and terminate a stalled parser."""
    def run():
        with tempfile.TemporaryDirectory(prefix='extract-', dir=private_directory) as directory:
            path = Path(directory) / 'input'
            path.write_bytes(data)
            path.chmod(0o600)
            child_env = {k: os.environ[k] for k in ('SYSTEMROOT', 'WINDIR', 'PATH') if k in os.environ}
            try:
                completed = subprocess.run([sys.executable, str(Path(__file__).resolve()), str(path), filename],
                    capture_output=True, timeout=60, env=child_env, check=False)
                if completed.returncode == 0:
                    return json.loads(completed.stdout)
                reason = 'document_parser_failed'
            except subprocess.TimeoutExpired:
                reason = 'document_parser_timeout'
            return {"segments": [], "gaps": [{"location": "file", "reason": reason}], "complete": False}
    return await asyncio.to_thread(run)


async def validate_download_url(url: str, canvas_url: str) -> None:
    parsed, canvas = urlsplit(url), urlsplit(canvas_url)
    if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError('invalid_attachment_location')
    if (parsed.scheme, parsed.hostname, parsed.port) == (canvas.scheme, canvas.hostname, canvas.port):
        return
    if parsed.scheme != 'https':
        raise ValueError('insecure_attachment_location')
    addresses = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
        raise ValueError('private_attachment_location')


async def download_attachment(attachment: dict, canvas_url: str, access_token: str) -> bytes:
    """Only call with an attachment returned by Canvas, never a student-authored URL."""
    if not attachment.get('url'):
        raise ValueError('attachment_download_unavailable')
    if (attachment.get('size') or 0) > MAX_FILE_BYTES:
        raise ValueError('file_size_limit')
    url = attachment['url']
    async with httpx2.AsyncClient(timeout=30, follow_redirects=False) as client:
        for _ in range(6):
            await validate_download_url(url, canvas_url)
            p, c = urlsplit(url), urlsplit(canvas_url)
            same_origin = (p.scheme, p.hostname, p.port) == (c.scheme, c.hostname, c.port)
            headers = {'Authorization': f'Bearer {access_token}'} if same_origin else {}
            async with client.stream('GET', url, headers=headers) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    location = response.headers.get('Location')
                    if not location:
                        raise ValueError('attachment_download_failed')
                    url = urljoin(url, location)
                    continue
                if response.status_code != 200:
                    raise ValueError('attachment_download_failed')
                chunks, total = [], 0
                async for chunk in response.aiter_bytes():
                    total += len(chunk)
                    if total > MAX_FILE_BYTES:
                        raise ValueError('file_size_limit')
                    chunks.append(chunk)
                return b''.join(chunks)
    raise ValueError('attachment_redirect_limit')


if __name__ == '__main__':
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import logging
    logging.disable(logging.CRITICAL)
    print(json.dumps(extract_document(Path(sys.argv[1]).read_bytes(), sys.argv[2]), ensure_ascii=False))
