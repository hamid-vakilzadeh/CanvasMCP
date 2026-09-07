import asyncio
import io
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from reporting.documents import extract_document, extract_isolated, download_attachment, validate_download_url


class DocumentTests(unittest.IsolatedAsyncioTestCase):
    def test_word_text_table_and_document_locations(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Synthetic reasoning supported by evidence.')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Claim'
        table.cell(0, 1).text = 'Evidence'
        data = io.BytesIO()
        doc.save(data)
        result = extract_document(data.getvalue(), 'work.docx')
        self.assertTrue(result['complete'])
        self.assertIn('Synthetic reasoning', result['segments'][0]['text'])
        self.assertTrue(any('row 1 cell 2' in s['location'] for s in result['segments']))

    def test_slides_and_notes(self):
        from pptx import Presentation
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[1])
        slide.shapes.title.text = 'Synthetic conclusion'
        slide.notes_slide.notes_text_frame.text = 'Support this conclusion with the sample.'
        data = io.BytesIO()
        doc.save(data)
        result = extract_document(data.getvalue(), 'work.pptx')
        self.assertTrue(any(s['location'] == 'slide 1 speaker notes' for s in result['segments']))
        self.assertIn('Synthetic conclusion', str(result))

    def test_workbook_preserves_formula_zero_and_missing_cached_result(self):
        from openpyxl import Workbook
        doc = Workbook()
        sheet = doc.active
        sheet.title = 'Analysis'
        sheet['A1'] = 0
        sheet['A2'] = 5
        sheet['A3'] = '=A1+A2'
        data = io.BytesIO()
        doc.save(data)
        result = extract_document(data.getvalue(), 'work.xlsx')
        self.assertIn({'location': 'sheet Analysis, cell A1', 'text': '0'}, result['segments'])
        self.assertIn('Formula: =A1+A2', str(result))
        self.assertIn('formula_result_unavailable_not_calculated', str(result))
        self.assertFalse(result['complete'])

    def test_scanned_pdf_and_unsupported_media_are_not_claimed_read(self):
        from pypdf import PdfWriter
        doc = PdfWriter()
        doc.add_blank_page(width=200, height=200)
        data = io.BytesIO()
        doc.write(data)
        self.assertFalse(extract_document(data.getvalue(), 'scan.pdf')['complete'])
        self.assertEqual(extract_document(b'video', 'work.mp4')['gaps'][0]['reason'], 'unsupported_format')

    async def test_isolated_csv_extraction_preserves_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            result = await extract_isolated(b'Item,Amount\nSynthetic,0\n', 'work.csv', Path(directory))
        self.assertTrue(result['complete'])
        self.assertEqual(result['segments'][1], {'location': 'row 2', 'text': 'Synthetic | 0'})

    async def test_attachment_redirect_does_not_forward_canvas_credentials(self):
        import httpx2
        seen = []
        def handler(request):
            seen.append((str(request.url), request.headers.get('authorization')))
            if request.url.host == 'canvas.example.invalid':
                return httpx2.Response(302, headers={'Location':'https://storage.example.invalid/work.txt'})
            return httpx2.Response(200, content=b'Synthetic file')
        original = httpx2.AsyncClient
        with patch('reporting.documents.httpx2.AsyncClient', lambda **kwargs: original(transport=httpx2.MockTransport(handler), **kwargs)), \
             patch('reporting.documents.socket.getaddrinfo', return_value=[(2,1,6,'',('93.184.216.34',443))]):
            data = await download_attachment({'url':'https://canvas.example.invalid/files/1/download'},
                                             'https://canvas.example.invalid','synthetic-secret')
        self.assertEqual(data,b'Synthetic file')
        self.assertEqual(seen[0][1],'Bearer synthetic-secret')
        self.assertIsNone(seen[1][1])
        with patch('reporting.documents.socket.getaddrinfo', return_value=[(2,1,6,'',('127.0.0.1',443))]):
            with self.assertRaises(ValueError):
                await validate_download_url('https://private.example.invalid/work','https://canvas.example.invalid')

    def test_encrypted_pdf_and_zip_expansion_limits_are_explicit(self):
        from pypdf import PdfWriter
        import zipfile
        doc=PdfWriter(); doc.add_blank_page(width=200,height=200); doc.encrypt('unavailable-password')
        data=io.BytesIO(); doc.write(data)
        self.assertIn('encrypted_document',str(extract_document(data.getvalue(),'locked.pdf')))
        data=io.BytesIO()
        with zipfile.ZipFile(data,'w',zipfile.ZIP_DEFLATED) as archive: archive.writestr('word/document.xml','x'*500)
        with patch('reporting.documents.MAX_EXPANDED_BYTES',100):
            self.assertIn('document_expansion_limit',str(extract_document(data.getvalue(),'oversized.docx')))

    def test_text_pdf_and_visual_html_coverage(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer=PdfWriter(); page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 240 Td (Synthetic PDF evidence) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
        data=io.BytesIO(); writer.write(data)
        result=extract_document(data.getvalue(),'text.pdf')
        self.assertTrue(result['complete'])
        self.assertEqual(result['segments'][0]['location'],'page 1')
        self.assertIn('Synthetic PDF evidence',result['segments'][0]['text'])
        visual=extract_document(b'<p>Visible text</p><img src="diagram.png">','work.html')
        self.assertFalse(visual['complete'])
        self.assertIn('Visible text',visual['segments'][0]['text'])


if __name__ == '__main__':
    unittest.main()
